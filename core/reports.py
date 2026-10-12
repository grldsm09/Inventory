from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def _excel_value(value: Any) -> Any:
    if value is None:
        return None
    try:
        missing = pd.isna(value)
        if getattr(missing, "ndim", 0) == 0 and bool(missing):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(value, "item") and callable(value.item):
        try:
            value = value.item()
        except (ValueError, TypeError):
            pass
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def _write_report(path: Path, sheet_name: str, rows: Iterable[dict[str, Any]]) -> Path:
    records = list(rows)
    headers = list(records[0]) if records else []
    if not headers:
        headers = {
            "Missing_Locations.xlsx": [
                "Main Location",
                "Specific Location",
                "Source Row",
                "Issue",
            ],
            "Duplicate_Assets.xlsx": ["Asset Tag No", "Duplicate Reason"],
            "Processing_Log.xlsx": ["Timestamp", "Level", "Message"],
        }.get(path.name, ["Status"])

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(headers)
    for record in records:
        worksheet.append([_excel_value(record.get(header)) for header in headers])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(records) + 1)}"
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243B53")
    for column in worksheet.columns:
        width = min(
            60,
            max(12, max(len(str(cell.value or "")) for cell in column) + 2),
        )
        worksheet.column_dimensions[column[0].column_letter].width = width
    if records:
        table = Table(
            displayName=f"{path.stem.replace('_', '')}Table",
            ref=f"A1:{get_column_letter(len(headers))}{len(records) + 1}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def export_reports(
    output_dir: str | Path,
    missing_locations: Iterable[dict[str, Any]],
    duplicate_records: pd.DataFrame,
    log_records: Iterable[dict[str, Any]],
) -> dict[str, Path]:
    directory = Path(output_dir)
    duplicate_rows = duplicate_records.to_dict(orient="records")
    return {
        "missing": _write_report(
            directory / "Missing_Locations.xlsx",
            "Missing Locations",
            missing_locations,
        ),
        "duplicates": _write_report(
            directory / "Duplicate_Assets.xlsx",
            "Duplicate Assets",
            duplicate_rows,
        ),
        "log": _write_report(
            directory / "Processing_Log.xlsx",
            "Processing Log",
            log_records,
        ),
    }


def export_processing_log(
    output_dir: str | Path,
    log_records: Iterable[dict[str, Any]],
) -> Path:
    return _write_report(
        Path(output_dir) / "Processing_Log.xlsx",
        "Processing Log",
        log_records,
    )


def make_log_record(message: str, level: str = "INFO") -> dict[str, str]:
    return {
        "Timestamp": datetime.now().astimezone().isoformat(timespec="seconds"),
        "Level": level,
        "Message": message,
    }
