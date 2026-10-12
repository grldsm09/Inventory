from __future__ import annotations

import copy
import hashlib
import math
import re
import shutil
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table
from openpyxl.utils import get_column_letter

from core.mapper import (
    LocationMapper,
    normalize_header,
    normalize_location,
    normalize_worksheet_key,
    resolve_column,
)
from core.pivot_refresh import refresh_pivots
from core.reports import export_reports, make_log_record
from core.settings import (
    DEFAULT_OUTPUT_DIR,
    LISTING_DATA_ROW,
    LISTING_HEADER_ROW,
    RAW_DATA_ROW,
    RAW_HEADER_ROW,
    ROOM_DATA_ROW,
    ROOM_HEADER_ROW,
    SPECIAL_SHEET_KEYS,
)
from core.validator import (
    duplicate_assets,
    filter_by_main_location,
    report_rows_for_missing,
    validate_filtered_records,
    validate_location_coverage,
    ValidationError,
)

ProgressCallback = Callable[[int, str], None]
LogCallback = Callable[[str], None]
PivotRefresher = Callable[[str | Path], None]


class WorkbookGenerationError(RuntimeError):
    def __init__(self, message: str, missing_locations: Iterable[str] = ()) -> None:
        super().__init__(message)
        self.missing_locations = list(missing_locations)


@dataclass
class GenerationResult:
    workbook_path: Path
    asset_count: int
    room_count: int
    duplicate_count: int
    missing_locations: list[str]
    report_paths: dict[str, Path]
    log_records: list[dict[str, str]]


def _to_excel_value(value: Any) -> Any:
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item") and callable(value.item):
        try:
            value = value.item()
        except (ValueError, TypeError):
            pass
    if isinstance(value, float) and math.isnan(value):
        return None
    if value is pd.NA or value is pd.NaT:
        return None
    return value


def _special_sheets(workbook: Any) -> dict[str, Any]:
    by_key = {
        normalize_worksheet_key(worksheet.title): worksheet
        for worksheet in workbook.worksheets
    }
    result: dict[str, Any] = {}
    for key, canonical_name in SPECIAL_SHEET_KEYS.items():
        sheet = by_key.get(normalize_worksheet_key(canonical_name))
        if sheet is not None:
            result[key] = sheet
    required = {"raw", "listingfile", "summarylisting"}
    missing = required - result.keys()
    if missing:
        raise WorkbookGenerationError(
            "Template is missing required worksheet(s): "
            + ", ".join(SPECIAL_SHEET_KEYS[key] for key in sorted(missing))
        )
    return result


def _find_room_template(workbook: Any, special_sheets: Mapping[str, Any], locations: list[str]) -> Any:
    special_ids = {id(sheet) for sheet in special_sheets.values()}
    wanted = {normalize_worksheet_key(location) for location in locations}
    for worksheet in workbook.worksheets:
        if id(worksheet) not in special_ids and normalize_worksheet_key(worksheet.title) in wanted:
            return worksheet
    for worksheet in workbook.worksheets:
        if id(worksheet) not in special_ids:
            return worksheet
    raise WorkbookGenerationError("Template does not contain a room worksheet to use for formatting.")


def _headers(worksheet: Any, header_row: int) -> list[tuple[int, str]]:
    return [
        (column, str(worksheet.cell(header_row, column).value).strip())
        for column in range(1, worksheet.max_column + 1)
        if worksheet.cell(header_row, column).value is not None
    ]


def _source_column(dataframe: pd.DataFrame, destination_header: str) -> object | None:
    header_key = normalize_header(destination_header)
    for source_column in dataframe.columns:
        if normalize_header(source_column) == header_key:
            return source_column

    for field in (
        "asset_tag",
        "asset_description",
        "quantity",
        "supplier",
        "acquisition_date",
        "unit_of_measure",
        "account_charge",
        "asset_category",
        "type_of_asset",
        "classification",
        "date_tagged",
        "unit_cost",
        "total_cost",
        "rr_number",
        "po_number",
        "serial_number",
        "main_location",
        "sub_location",
        "specific_location",
        "department",
        "rc_code",
        "program",
        "pc_code",
        "custodian",
        "remarks",
    ):
        if resolve_column([destination_header], field) is not None:
            return resolve_column(dataframe.columns, field)
    return None


def _clear_data_rows(worksheet: Any, header_row: int) -> None:
    first_data = header_row + 1
    if worksheet.max_row >= first_data:
        worksheet.delete_rows(first_data, worksheet.max_row - first_data + 1)


def _set_filter(worksheet: Any, header_row: int, row_count: int, column_count: int) -> None:
    if column_count:
        last_column = get_column_letter(column_count)
        last_row = max(header_row, header_row + row_count)
        worksheet.auto_filter.ref = f"A{header_row}:{last_column}{last_row}"


def _append_records(
    worksheet: Any,
    dataframe: pd.DataFrame,
    header_row: int,
    first_data_row: int,
    count_column: bool = False,
) -> None:
    columns = _headers(worksheet, header_row)
    source_for_column = {
        column_index: _source_column(dataframe, heading)
        for column_index, heading in columns
    }
    if count_column:
        for column_index, heading in columns:
            if normalize_header(heading) == "count":
                source_for_column[column_index] = None

    template_row = first_data_row if worksheet.max_row >= first_data_row else None
    row_number = first_data_row
    for count, (_, record) in enumerate(dataframe.iterrows(), start=1):
        for column_index, heading in columns:
            cell = worksheet.cell(row=row_number, column=column_index)
            if template_row is not None and row_number != template_row:
                source_cell = worksheet.cell(row=template_row, column=column_index)
                if source_cell.has_style:
                    cell._style = copy.copy(source_cell._style)
                if source_cell.number_format:
                    cell.number_format = source_cell.number_format
            if count_column and normalize_header(heading) == "count":
                cell.value = count
                continue
            source_column = source_for_column[column_index]
            cell.value = (
                _to_excel_value(record.get(source_column))
                if source_column is not None
                else None
            )
        row_number += 1

    _set_filter(worksheet, header_row, len(dataframe), worksheet.max_column)


def _table_names(workbook: Any) -> set[str]:
    return {
        table.displayName.casefold()
        for worksheet in workbook.worksheets
        for table in worksheet.tables.values()
    }


def _resize_tables(worksheet: Any, header_row: int, row_count: int) -> None:
    if not worksheet.tables:
        return
    for table in worksheet.tables.values():
        min_col = worksheet[table.ref.split(":")[0]].column
        max_col = worksheet[table.ref.split(":")[-1]].column
        table.ref = (
            f"{get_column_letter(min_col)}{header_row}:"
            f"{get_column_letter(max_col)}{max(header_row, header_row + row_count)}"
        )
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref


def _unique_table_name(location_key: str, used: set[str]) -> str:
    digest = hashlib.sha1(location_key.encode("utf-8")).hexdigest()[:10]
    base = f"Room_{digest}"
    candidate = base
    suffix = 2
    while candidate.casefold() in used:
        candidate = f"{base}_{suffix}"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _attach_room_tables(
    worksheet: Any,
    table_templates: list[Table],
    location_key: str,
    row_count: int,
    used_names: set[str],
) -> None:
    for name in list(worksheet.tables.keys()):
        del worksheet.tables[name]
    for source_table in table_templates:
        table = copy.deepcopy(source_table)
        table.displayName = _unique_table_name(location_key, used_names)
        table.name = table.displayName
        first_cell, _ = source_table.ref.split(":")
        column_match = re.match(r"[A-Z]+", first_cell)
        if column_match is None:
            raise WorkbookGenerationError(
                f"Invalid room-table reference in template: {source_table.ref}"
            )
        first_column = column_match.group(0)
        final_column = get_column_letter(worksheet.max_column)
        table.ref = (
            f"{first_column}{ROOM_HEADER_ROW}:"
            f"{final_column}{max(ROOM_HEADER_ROW, ROOM_HEADER_ROW + row_count)}"
        )
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
        worksheet.add_table(table)


def _trim_room_template(worksheet: Any, tables: list[Table]) -> None:
    if worksheet.max_row > ROOM_DATA_ROW:
        worksheet.delete_rows(ROOM_DATA_ROW + 1, worksheet.max_row - ROOM_DATA_ROW)
    for table in tables:
        table.ref = (
            f"{table.ref.split(':')[0]}:{get_column_letter(worksheet.max_column)}"
            f"{ROOM_DATA_ROW}"
        )
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref


def _clear_room_data_values(worksheet: Any) -> None:
    for row in worksheet.iter_rows(
        min_row=ROOM_DATA_ROW,
        max_row=worksheet.max_row,
        max_col=worksheet.max_column,
    ):
        for cell in row:
            cell.value = None


def _build_log(callback: LogCallback | None, entries: list[dict[str, str]], message: str) -> None:
    entries.append(make_log_record(message))
    if callback:
        callback(message)


def generate_workbook(
    dataframe: pd.DataFrame,
    selected_location: str,
    template_path: str | Path,
    output_dir: str | Path = DEFAULT_OUTPUT_DIR,
    main_location_column: object | None = None,
    specific_location_column: object | None = None,
    location_overrides: Mapping[str, str] | None = None,
    progress_callback: ProgressCallback | None = None,
    log_callback: LogCallback | None = None,
    pivot_refresher: PivotRefresher = refresh_pivots,
) -> GenerationResult:
    template = Path(template_path).expanduser().resolve()
    directory = Path(output_dir).expanduser().resolve()
    if not template.is_file():
        raise WorkbookGenerationError(f"Template workbook does not exist: {template}")

    main_column = (
        main_location_column
        if main_location_column is not None
        else resolve_column(dataframe.columns, "main_location")
    )
    room_column = (
        specific_location_column
        if specific_location_column is not None
        else resolve_column(dataframe.columns, "specific_location")
    )
    if main_column is None or room_column is None:
        raise WorkbookGenerationError(
            "The source data must contain Main Location and Specific Location columns."
        )

    filtered = filter_by_main_location(dataframe, main_column, selected_location)
    room_values = filtered[room_column].tolist()
    mapper = LocationMapper(room_values, location_overrides)
    try:
        missing_before = validate_filtered_records(
            filtered,
            main_column,
            room_column,
            selected_location,
            mapper,
        )
    except ValidationError as exc:
        directory.mkdir(parents=True, exist_ok=True)
        invalid_room_rows = report_rows_for_missing(
            filtered,
            room_column,
            selected_location,
            exc.missing_locations,
        )
        export_reports(directory, invalid_room_rows, duplicate_assets(filtered), [])
        raise WorkbookGenerationError(str(exc), exc.missing_locations) from exc
    if missing_before:
        raise WorkbookGenerationError(
            "Location mapping is incomplete: " + ", ".join(missing_before),
            missing_before,
        )

    canonical_column = "_InventoryGeneratorCanonicalRoom"
    filtered[canonical_column] = filtered[room_column].map(mapper.resolve)
    unique_rooms = mapper.locations
    worksheet_by_key = mapper.worksheet_names(SPECIAL_SHEET_KEYS.values())
    if progress_callback:
        progress_callback(5, f"Preparing {len(filtered):,} assets across {len(unique_rooms):,} rooms")

    directory.mkdir(parents=True, exist_ok=True)
    safe_building = re.sub(r'[<>:"/\\|?*]+', "-", selected_location).strip().strip(". ")
    if not safe_building:
        safe_building = "Building"
    output_path = directory / f"Inventory Location and Listings_{safe_building}.xlsx"
    temporary_path = directory / f"{output_path.stem}_BUILDING_{uuid.uuid4().hex}.xlsx"
    try:
        shutil.copy2(template, temporary_path)
    except OSError as exc:
        temporary_path.unlink(missing_ok=True)
        raise WorkbookGenerationError(
            f"Could not create a working copy of the template: {exc}"
        ) from exc

    log_records: list[dict[str, str]] = []
    _build_log(log_callback, log_records, f"Copied template: {template}")

    workbook: Workbook | None = None
    try:
        workbook = load_workbook(temporary_path)
        special = _special_sheets(workbook)
        location_count_sheet = special.get("locationcountperfile")
        if location_count_sheet is None:
            location_count_sheet = workbook.create_sheet(
                SPECIAL_SHEET_KEYS["locationcountperfile"]
            )
            special["locationcountperfile"] = location_count_sheet

        room_template = _find_room_template(workbook, special, unique_rooms)
        room_tables = [copy.deepcopy(table) for table in room_template.tables.values()]
        if not room_tables:
            raise WorkbookGenerationError(
                f"Template room worksheet '{room_template.title}' does not contain a formatting table."
            )

        used_tables = _table_names(workbook)
        _trim_room_template(room_template, room_tables)
        format_seed = workbook.copy_worksheet(room_template)
        format_seed.title = "__RoomFormatSeed__"
        workbook.remove(room_template)

        for worksheet in list(workbook.worksheets):
            if worksheet is format_seed or any(worksheet is item for item in special.values()):
                continue
            workbook.remove(worksheet)

        _clear_data_rows(special["raw"], RAW_HEADER_ROW)
        _append_records(
            special["raw"],
            filtered.drop(columns=[canonical_column]),
            RAW_HEADER_ROW,
            RAW_DATA_ROW,
        )
        _resize_tables(special["raw"], RAW_HEADER_ROW, len(filtered))

        listing = special["listingfile"]
        _clear_data_rows(listing, LISTING_HEADER_ROW)
        _append_records(
            listing,
            filtered.drop(columns=[canonical_column]),
            LISTING_HEADER_ROW,
            LISTING_DATA_ROW,
        )
        _resize_tables(listing, LISTING_HEADER_ROW, len(filtered))

        created_by_key: dict[str, str] = {}
        for index, location in enumerate(unique_rooms, start=1):
            location_key = normalize_location(location)
            sheet_title = worksheet_by_key[location_key]
            room_sheet = workbook.copy_worksheet(format_seed)
            room_sheet.title = sheet_title
            title_cell = room_sheet["A1"]
            if isinstance(title_cell.value, str) and ":" in title_cell.value:
                title_cell.value = title_cell.value.rsplit(":", 1)[0] + f": {location}"

            records = filtered.loc[filtered[canonical_column].map(normalize_location).eq(location_key)]
            _clear_room_data_values(room_sheet)
            _append_records(
                room_sheet,
                records.drop(columns=[canonical_column]),
                ROOM_HEADER_ROW,
                ROOM_DATA_ROW,
                count_column=True,
            )
            _attach_room_tables(
                room_sheet,
                room_tables,
                location_key,
                len(records),
                used_tables,
            )
            created_by_key[location_key] = room_sheet.title
            if progress_callback:
                progress = 10 + int(index * 70 / max(1, len(unique_rooms)))
                progress_callback(progress, f"Building worksheet {index}/{len(unique_rooms)}: {location}")

        missing = validate_location_coverage(mapper.location_keys, created_by_key)
        if missing:
            missing_rows = report_rows_for_missing(
                filtered,
                room_column,
                selected_location,
                missing,
            )
            export_reports(directory, missing_rows, duplicate_assets(filtered), log_records)
            raise WorkbookGenerationError(
                "Workbook validation failed. No workbook was saved because worksheets "
                "are missing for: " + ", ".join(missing),
                missing,
            )

        workbook.remove(format_seed)

        allowed_special_keys = set(SPECIAL_SHEET_KEYS)
        expected_sheet_keys = {
            normalize_worksheet_key(SPECIAL_SHEET_KEYS[key])
            for key in allowed_special_keys
            if key in special
        }
        actual_special_keys = {
            normalize_worksheet_key(sheet.title)
            for sheet in workbook.worksheets
            if normalize_worksheet_key(sheet.title) in expected_sheet_keys
        }
        unexpected = [
            sheet.title
            for sheet in workbook.worksheets
            if normalize_worksheet_key(sheet.title) not in expected_sheet_keys
            and sheet.title not in created_by_key.values()
        ]
        if unexpected or actual_special_keys != expected_sheet_keys:
            raise WorkbookGenerationError(
                "Workbook sheet validation failed. Unexpected sheets: "
                + (", ".join(unexpected) if unexpected else "none")
            )

        workbook.save(temporary_path)
        workbook.close()
        workbook = None
    except WorkbookGenerationError:
        if workbook is not None:
            workbook.close()
        temporary_path.unlink(missing_ok=True)
        raise
    except Exception as exc:
        if workbook is not None:
            workbook.close()
        temporary_path.unlink(missing_ok=True)
        raise WorkbookGenerationError(
            f"Could not build workbook '{output_path}': {exc}"
        ) from exc

    _build_log(
        log_callback,
        log_records,
        f"Saved workbook data with {len(filtered):,} records in {len(unique_rooms):,} location sheets",
    )
    if progress_callback:
        progress_callback(85, "Refreshing Excel PivotTables...")

    duplicates = duplicate_assets(filtered.drop(columns=[canonical_column]))
    missing_report_rows: list[dict[str, Any]] = []
    try:
        pivot_refresher(temporary_path)
    except Exception as exc:
        _build_log(log_callback, log_records, f"Pivot refresh failed: {exc}")
        try:
            export_reports(directory, missing_report_rows, duplicates, log_records)
        finally:
            temporary_path.unlink(missing_ok=True)
        raise WorkbookGenerationError(
            f"Workbook data was generated but PivotTables were not refreshed. "
            f"The workbook was not marked complete. {exc}"
        ) from exc

    _build_log(log_callback, log_records, "PivotTables refreshed and workbook saved by Excel")
    if progress_callback:
        progress_callback(95, "Writing processing reports...")
    _build_log(log_callback, log_records, "Writing processing reports")
    try:
        reports = export_reports(directory, missing_report_rows, duplicates, log_records)
        temporary_path.replace(output_path)
    except Exception as exc:
        temporary_path.unlink(missing_ok=True)
        raise WorkbookGenerationError(
            f"Could not write processing reports or finalize '{output_path}': {exc}"
        ) from exc
    if progress_callback:
        progress_callback(100, "Workbook generation complete")

    return GenerationResult(
        workbook_path=output_path,
        asset_count=len(filtered),
        room_count=len(unique_rooms),
        duplicate_count=len(duplicates),
        missing_locations=[],
        report_paths=reports,
        log_records=log_records,
    )
